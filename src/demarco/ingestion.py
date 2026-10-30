"""Leitura sem escrita nos originais, com origem e tipos das células preservados."""

import hashlib
import json
from dataclasses import dataclass
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from demarco.sources import Source


class InputError(ValueError):
    """A entrada não atende ao contrato; a execução deve parar antes de exportar."""


def json_default(value: Any) -> str:
    if isinstance(value, (date, datetime, time)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    raise TypeError(f"Tipo não serializável: {type(value).__name__}")


def encode_json(value: Any) -> str:
    return json.dumps(
        value, ensure_ascii=False, sort_keys=True, default=json_default, allow_nan=False
    )


@dataclass
class Record:
    row: int
    values: dict[str, Any]


@dataclass
class Table:
    source: Source
    sheet: str
    sha256: str
    records: list[Record]
    skipped_blank_rows: int


def read_source(directory: Path, source: Source) -> Table:
    path = directory / source.filename
    if not path.is_file():
        raise InputError(f"Arquivo não encontrado: {path}")
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    # data_only=False: fórmulas não podem virar nulos silenciosos por falta de cache.
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        if len(workbook.worksheets) != 1:
            raise InputError(f"{source.filename}: esperada uma única aba de dados.")
        sheet = workbook.worksheets[0]
        iterator = sheet.iter_rows()
        first = next(iterator, ())
        headers = tuple(cell.value for cell in first)
        if len(headers) != len(set(headers)):
            raise InputError(f"{source.filename}: cabeçalhos duplicados.")
        if set(headers) != set(source.columns):
            raise InputError(
                f"{source.filename}: cabeçalhos incompatíveis. "
                f"Esperados: {source.columns}; recebidos: {headers}"
            )
        records = []
        skipped = 0
        for row_number, cells in enumerate(iterator, start=2):
            if all(cell.value is None for cell in cells):
                skipped += 1
                continue
            for cell in cells:
                if cell.data_type in {"f", "e"}:
                    raise InputError(
                        f"{source.filename}, {cell.coordinate}: fórmula ou erro Excel. "
                        "Forneça valores explícitos antes da ingestão."
                    )
            values = [cell.value for cell in cells]
            if any(value is not None for value in values[len(headers) :]):
                raise InputError(
                    f"{source.filename}, linha {row_number}: dados em coluna sem cabeçalho."
                )
            # Exportações do Sheets podem omitir a dimensão da aba e as células
            # vazias ao fim de cada linha. Completar só esses vazios, sem deslocar valores.
            values = values[: len(headers)] + [None] * max(0, len(headers) - len(values))
            records.append(Record(row_number, dict(zip(headers, values, strict=True))))
        return Table(source, sheet.title, digest, records, skipped)
    finally:
        workbook.close()


def export_table(table: Table, path: Path) -> None:
    with path.open("w", encoding="utf-8") as stream:
        for record in table.records:
            stream.write(
                encode_json(
                    {
                        "source_file": table.source.filename,
                        "source_sheet": table.sheet,
                        "source_row": record.row,
                        "values": record.values,
                        "value_types": {k: type(v).__name__ for k, v in record.values.items()},
                    }
                )
                + "\n"
            )
